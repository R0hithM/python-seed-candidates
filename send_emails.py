#!/usr/bin/env python3
"""Send plain-text emails from an Excel workbook using SMTP templates."""

from __future__ import annotations

import argparse
import os
import re
import smtplib
import sys
from dataclasses import dataclass
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

ERROR_BODY_MAX_LEN = 500
DEFAULT_EXCEL_FILE = "Emails.xlsx"
DEFAULT_TEMPLATE_DIR = "email_template"
DEFAULT_SMTP_PORT = 587
SMTP_SSL_PORT = 465

EMAIL_HEADER = "Email"
TEMPLATE_NAME_HEADER = "Template Name"

COL_EMAIL_STATUS = "email_status"
COL_EMAIL_ERROR = "email_error"
COL_EMAIL_SENT_AT = "email_sent_at"

RESULT_COLUMNS = (
    COL_EMAIL_STATUS,
    COL_EMAIL_ERROR,
    COL_EMAIL_SENT_AT,
)

STATUS_SUCCESS = "SUCCESS"
STATUS_FAILED = "FAILED"
STATUS_SKIPPED = "SKIPPED"
TERMINAL_STATUSES = {STATUS_SUCCESS, STATUS_FAILED, STATUS_SKIPPED}

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
SUBJECT_PREFIX = "subject:"


@dataclass(frozen=True)
class SmtpConfig:
    host: str
    port: int
    user: str
    password: str
    from_addr: str


def load_config() -> SmtpConfig:
    load_dotenv()
    host = os.getenv("SMTP_HOST", "").strip()
    user = os.getenv("SMTP_USER", "").strip()
    password = os.getenv("SMTP_PASSWORD", "").strip()
    from_addr = os.getenv("SMTP_FROM", "").strip() or user
    port_raw = os.getenv("SMTP_PORT", str(DEFAULT_SMTP_PORT)).strip() or str(
        DEFAULT_SMTP_PORT
    )
    try:
        port = int(port_raw)
    except ValueError:
        print(f"SMTP_PORT must be an integer, got {port_raw!r}.", file=sys.stderr)
        sys.exit(1)

    missing = [
        name
        for name, value in (
            ("SMTP_HOST", host),
            ("SMTP_USER", user),
            ("SMTP_PASSWORD", password),
        )
        if not value
    ]
    if missing:
        print(
            "Missing SMTP settings: "
            + ", ".join(missing)
            + ". Copy .env.example to .env and set them.",
            file=sys.stderr,
        )
        sys.exit(1)
    if not EMAIL_RE.match(from_addr):
        print(f"SMTP_FROM / SMTP_USER is not a valid email: {from_addr}", file=sys.stderr)
        sys.exit(1)
    return SmtpConfig(
        host=host,
        port=port,
        user=user,
        password=password,
        from_addr=from_addr,
    )


def header_map(ws: Worksheet) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for cell in ws[1]:
        if cell.value is None:
            continue
        mapping[str(cell.value).strip()] = cell.column
    return mapping


def ensure_result_columns(ws: Worksheet, headers: dict[str, int]) -> dict[str, int]:
    next_col = ws.max_column
    for name in RESULT_COLUMNS:
        if name in headers:
            continue
        next_col += 1
        ws.cell(row=1, column=next_col, value=name)
        headers[name] = next_col
    return headers


def cell_str(ws: Worksheet, row: int, col: int | None) -> str:
    if col is None:
        return ""
    value = ws.cell(row=row, column=col).value
    if value is None:
        return ""
    return str(value).strip()


def iso_now() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def truncate_error(text: str) -> str:
    text = " ".join(text.split())
    if len(text) <= ERROR_BODY_MAX_LEN:
        return text
    return text[: ERROR_BODY_MAX_LEN - 3] + "..."


def write_result(
    ws: Worksheet,
    row: int,
    headers: dict[str, int],
    *,
    status: str,
    error: str = "",
    sent_at: str = "",
) -> None:
    ws.cell(row=row, column=headers[COL_EMAIL_STATUS], value=status)
    ws.cell(row=row, column=headers[COL_EMAIL_ERROR], value=error)
    ws.cell(row=row, column=headers[COL_EMAIL_SENT_AT], value=sent_at)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Send plain-text emails from an Excel file using SMTP templates."
    )
    parser.add_argument(
        "--file",
        default=DEFAULT_EXCEL_FILE,
        help=f"Path to the Excel workbook (default: {DEFAULT_EXCEL_FILE})",
    )
    parser.add_argument(
        "--templates",
        default=DEFAULT_TEMPLATE_DIR,
        help=f"Folder of .txt templates (default: {DEFAULT_TEMPLATE_DIR})",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=None,
        help="Process at most N pending rows this run (for smoke tests)",
    )
    parser.add_argument(
        "--retry-failed",
        action="store_true",
        help="Re-attempt rows previously marked FAILED (SUCCESS and SKIPPED still skipped)",
    )
    return parser.parse_args()


def should_skip_row(status: str, retry_failed: bool) -> bool:
    if not status:
        return False
    if retry_failed and status == STATUS_FAILED:
        return False
    return status in TERMINAL_STATUSES


def is_valid_email(address: str) -> bool:
    return bool(EMAIL_RE.match(address))


def template_path(template_dir: Path, template_name: str) -> Path | None:
    safe_name = Path(template_name).name
    if not safe_name or safe_name in {".", ".."}:
        return None
    if not safe_name.lower().endswith(".txt"):
        safe_name = f"{safe_name}.txt"
    path = (template_dir / safe_name).resolve()
    try:
        path.relative_to(template_dir.resolve())
    except ValueError:
        return None
    return path


def parse_template(text: str) -> tuple[str, str] | str:
    """Return (subject, body) or an error string."""
    lines = text.splitlines()
    if not lines:
        return "template is empty"
    first = lines[0].strip()
    if not first.lower().startswith(SUBJECT_PREFIX):
        return "first line must start with 'Subject:'"
    subject = first[len(SUBJECT_PREFIX) :].strip()
    if not subject:
        return "Subject: line is empty"
    body = "\n".join(lines[1:]).lstrip("\n")
    return subject, body


def send_email(
    config: SmtpConfig,
    to_addr: str,
    subject: str,
    body: str,
) -> str | None:
    """Send a plain-text email. Return error text, or None on success."""
    message = EmailMessage()
    message["From"] = config.from_addr
    message["To"] = to_addr
    message["Subject"] = subject
    message.set_content(body, subtype="plain")

    try:
        if config.port == SMTP_SSL_PORT:
            with smtplib.SMTP_SSL(config.host, config.port, timeout=30) as smtp:
                smtp.login(config.user, config.password)
                smtp.send_message(message)
        else:
            with smtplib.SMTP(config.host, config.port, timeout=30) as smtp:
                smtp.ehlo()
                smtp.starttls()
                smtp.ehlo()
                smtp.login(config.user, config.password)
                smtp.send_message(message)
    except smtplib.SMTPRecipientsRefused as exc:
        return truncate_error(f"recipient refused: {exc}")
    except smtplib.SMTPAuthenticationError as exc:
        return truncate_error(f"SMTP authentication failed: {exc}")
    except smtplib.SMTPException as exc:
        return truncate_error(f"SMTP error: {exc}")
    except OSError as exc:
        return truncate_error(f"connection error: {exc}")
    return None


def main() -> int:
    args = parse_args()
    config = load_config()

    path = Path(args.file)
    if not path.exists():
        print(f"Excel file not found: {path}", file=sys.stderr)
        return 1

    template_dir = Path(args.templates)
    if not template_dir.is_dir():
        print(f"Templates folder not found: {template_dir}", file=sys.stderr)
        return 1

    workbook = load_workbook(path)
    ws = workbook.active
    headers = ensure_result_columns(ws, header_map(ws))

    email_col = headers.get(EMAIL_HEADER)
    template_col = headers.get(TEMPLATE_NAME_HEADER)
    missing_cols = [
        name
        for name, col in ((EMAIL_HEADER, email_col), (TEMPLATE_NAME_HEADER, template_col))
        if col is None
    ]
    if missing_cols:
        print(
            "Column(s) not found in "
            + f"{path}: {', '.join(repr(name) for name in missing_cols)}",
            file=sys.stderr,
        )
        return 1

    processed = 0
    success = 0
    failed = 0
    skipped = 0
    already_done = 0
    dirty = False

    def save() -> None:
        nonlocal dirty
        if not dirty:
            return
        workbook.save(path)
        dirty = False

    max_row = ws.max_row
    print(
        f"Loaded {path} sheet={ws.title!r} rows={max_row - 1} "
        f"smtp={config.host}:{config.port} templates={template_dir}"
    )

    try:
        for row in range(2, max_row + 1):
            if args.limit is not None and processed >= args.limit:
                break

            status = cell_str(ws, row, headers[COL_EMAIL_STATUS])
            if should_skip_row(status, args.retry_failed):
                already_done += 1
                continue

            to_addr = cell_str(ws, row, email_col)
            template_name = cell_str(ws, row, template_col)
            label = to_addr or f"row {row}"

            if not to_addr:
                write_result(
                    ws,
                    row,
                    headers,
                    status=STATUS_SKIPPED,
                    error="missing email",
                    sent_at=iso_now(),
                )
                dirty = True
                save()
                skipped += 1
                processed += 1
                print(f"[{row}/{max_row}] SKIPPED {label}: missing email")
                continue

            if not is_valid_email(to_addr):
                write_result(
                    ws,
                    row,
                    headers,
                    status=STATUS_FAILED,
                    error="invalid email",
                    sent_at=iso_now(),
                )
                dirty = True
                save()
                failed += 1
                processed += 1
                print(f"[{row}/{max_row}] FAILED {label}: invalid email")
                continue

            if not template_name:
                write_result(
                    ws,
                    row,
                    headers,
                    status=STATUS_SKIPPED,
                    error="missing template name",
                    sent_at=iso_now(),
                )
                dirty = True
                save()
                skipped += 1
                processed += 1
                print(f"[{row}/{max_row}] SKIPPED {label}: missing template name")
                continue

            tmpl_path = template_path(template_dir, template_name)
            if tmpl_path is None or not tmpl_path.is_file():
                reason = (
                    f"template {template_name!r} not found in email_template folder"
                )
                write_result(
                    ws,
                    row,
                    headers,
                    status=STATUS_FAILED,
                    error=reason,
                    sent_at=iso_now(),
                )
                dirty = True
                save()
                failed += 1
                processed += 1
                print(f"[{row}/{max_row}] FAILED {label}: {reason}")
                continue

            try:
                template_text = tmpl_path.read_text(encoding="utf-8")
            except OSError as exc:
                reason = truncate_error(f"could not read template: {exc}")
                write_result(
                    ws,
                    row,
                    headers,
                    status=STATUS_FAILED,
                    error=reason,
                    sent_at=iso_now(),
                )
                dirty = True
                save()
                failed += 1
                processed += 1
                print(f"[{row}/{max_row}] FAILED {label}: {reason}")
                continue

            parsed = parse_template(template_text)
            if isinstance(parsed, str):
                write_result(
                    ws,
                    row,
                    headers,
                    status=STATUS_FAILED,
                    error=parsed,
                    sent_at=iso_now(),
                )
                dirty = True
                save()
                failed += 1
                processed += 1
                print(f"[{row}/{max_row}] FAILED {label}: {parsed}")
                continue

            subject, body = parsed
            print(f"[{row}/{max_row}] SEND {label} template={tmpl_path.name} ...", flush=True)
            smtp_error = send_email(config, to_addr, subject, body)
            sent_at = iso_now()
            if smtp_error:
                write_result(
                    ws,
                    row,
                    headers,
                    status=STATUS_FAILED,
                    error=smtp_error,
                    sent_at=sent_at,
                )
                dirty = True
                save()
                failed += 1
                processed += 1
                print(f"[{row}/{max_row}] FAILED {label}: {smtp_error}")
                continue

            write_result(
                ws,
                row,
                headers,
                status=STATUS_SUCCESS,
                sent_at=sent_at,
            )
            dirty = True
            save()
            success += 1
            processed += 1
            print(f"[{row}/{max_row}] SUCCESS {label} sent_at={sent_at}")
    except KeyboardInterrupt:
        print("\nInterrupted; saving workbook...", file=sys.stderr)
        save()
        print(
            f"Stopped. processed={processed} success={success} failed={failed} "
            f"skipped={skipped} already_done={already_done}"
        )
        return 130
    finally:
        save()

    print(
        f"Done. processed={processed} success={success} failed={failed} "
        f"skipped={skipped} already_done={already_done}"
    )
    return 0 if failed == 0 else 2


if __name__ == "__main__":
    sys.exit(main())
